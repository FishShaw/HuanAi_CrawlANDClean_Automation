#!/usr/bin/env python3
"""把逐市查到的官网和环评公示入口，并到全省单位表上，导出 XLSX。

    python build_entries_xlsx.py                      # 默认江苏（320000），完整单位表
    python build_entries_xlsx.py --province 340000     # 换省
    python build_entries_xlsx.py --province 340000 --approvers   # 只出审批机构口径

输入：
    <省>环评单位.csv / <省>环评审批机构.csv   单位清单（merge_jiangsu.py 产出）
    entries/<市编码>_<市>.csv                逐市录入的入口，列见 COLUMNS_IN

输出：
    <省>环评单位_官网及公示入口.xlsx（或 <省>环评审批机构_官网及公示入口.xlsx）

四个入口列是固定的：没有分设栏目、三类公示混在同一个页面时，把那个混合入口
重复填进对应的几列，并在备注里注明「三类合并在同一栏目」——这样下游筛选不用判断该看哪列。

没录入的单位留空，核验情况写「未查」，方便看进度。
"""
from __future__ import annotations

import argparse
import csv
import json
import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ENTRIES_DIR = HERE / "entries"
RAW_DIR = HERE / "raw"

# entries/*.csv 的列
COLUMNS_IN = ["单位", "首页入口", "受理公示入口", "审批前公示入口", "审批后意见入口",
              "入口（点击路径）", "备注", "核验情况", "核验日期"]
# 输出表的列
COLUMNS_OUT = ["城市", "地区", "单位", "记录数", *COLUMNS_IN[1:]]
WIDTHS = [10, 14, 34, 8, 34, 40, 40, 40, 34, 46, 16, 12]

# 非环保部门：按约定只给主页，不逐个查环评栏目
ECO = re.compile(r"生态环境|环境保护|环保")
ADMIN = re.compile(r"行政审批|管委会|管理委员会|政务服务|行政服务|数据局|市民中心"
                   r"|开发区|高新区|新区|园区|示范区|保税区|港区")


def tier(name: str) -> str:
    if ECO.search(name):
        return "生态环境系统"
    if ADMIN.search(name):
        return "开发区/行政审批"
    return "其他部门"


def load_entries() -> dict[tuple[str, str], dict]:
    """按 (城市, 单位) 索引。城市取自文件名，如 320600_南通市.csv。

    不能只按单位名索引：「高新区行政审批局」在南通（海安高新区）和泰州（海陵高新区）
    各有一个，同名不同单位。000000_省级及国家.csv 对应城市列里的“—”。
    """
    out: dict[tuple[str, str], dict] = {}
    if not ENTRIES_DIR.exists():
        return out
    for f in sorted(ENTRIES_DIR.glob("*.csv")):
        city = f.stem.split("_", 1)[-1]
        if f.stem.startswith("000000"):
            city = "—"
        with open(f, encoding="utf-8-sig") as fh:
            for row in csv.DictReader(fh):
                name = (row.get("单位") or "").strip()
                if name:
                    out[(city, name)] = row
    return out


def province_name(code: str) -> str:
    """省名取自抓取时存下来的地区树，不写死地名。"""
    for f in sorted(RAW_DIR.glob(f"{code[:2]}*.json")):
        tree = json.loads(f.read_text(encoding="utf-8"))["地区树"]
        if tree["编码"] == code:
            return tree["名称"]
    raise SystemExit(f"raw/ 里没有省编码 {code} 的地区树，先抓数据或改 --province")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--province", default="320000", help="省编码，默认 320000（江苏）")
    ap.add_argument("--approvers", action="store_true",
                    help="用 <省>环评审批机构.csv（去掉其他部门的交付口径），默认用完整单位表")
    args = ap.parse_args()

    prov = province_name(args.province)
    kind = "环评审批机构" if args.approvers else "环评单位"
    units_path = HERE / f"{prov}{kind}.csv"
    if not units_path.exists():
        raise SystemExit(f"没有 {units_path.name}，先跑 merge_jiangsu.py --province {args.province}")
    out_path = HERE / f"{prov}{kind}_官网及公示入口.xlsx"
    units = list(csv.DictReader(open(units_path, encoding="utf-8-sig")))
    entries = load_entries()
    records = sum(int(u["记录数"]) for u in units)
    cities = len({u["城市"] for u in units} - {"—"})

    wb = Workbook()
    ws = wb.active
    ws.title = "官网及公示入口"

    ws["A1"] = f"{prov}{kind}：官网及环评公示入口"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (f"单位来自 {units_path.name}（{cities} 市、{records} 条环评受理记录清洗归并）。"
                f"四个入口列固定；没有分设栏目的，混合入口重复填在相关列并在备注注明。"
                f"「其他部门」按约定只给主页。生成于 {date.today().isoformat()}。")
    ws["A3"] = ("核验情况：已核实=打开页面确认栏目存在；未核实=按站点结构给出最可能入口，未逐条打开；"
                "访问受限=本次请求被拒或超时，不代表浏览器打不开；无独立环评栏目=该单位未设公示专栏。")
    for r in (2, 3):
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(COLUMNS_OUT))

    head = 5
    for j, name in enumerate(COLUMNS_OUT, 1):
        c = ws.cell(head, j, name)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS[j - 1]

    done = 0
    for i, u in enumerate(units, head + 1):
        e = entries.get((u["城市"], u["单位"]), {})
        if e:
            done += 1
        vals = [u["城市"], u["地区"], u["单位"], int(u["记录数"])]
        for col in COLUMNS_IN[1:]:
            vals.append(e.get(col, "") or ("" if e else ""))
        if not e:
            vals[COLUMNS_OUT.index("核验情况")] = "未查"
            vals[COLUMNS_OUT.index("备注")] = tier(u["单位"])
        for j, v in enumerate(vals, 1):
            cell = ws.cell(i, j, v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = ws.cell(head + 1, 4)
    ws.auto_filter.ref = f"A{head}:{get_column_letter(len(COLUMNS_OUT))}{head + len(units)}"
    wb.save(out_path)
    print(f"{len(units)} 个单位，已录入 {done} 个，待查 {len(units) - done} 个 → {out_path}")
    by_city: dict[str, list[int]] = {}
    for u in units:
        d = by_city.setdefault(u["城市"], [0, 0])
        d[0] += 1
        d[1] += 1 if (u["城市"], u["单位"]) in entries else 0
    for city, (n, k) in sorted(by_city.items()):
        print(f"  {city}: {k}/{n}")


if __name__ == "__main__":
    main()
